"""Assemble the mid-semester review presentation (PowerPoint .pptx).

The slide text is hand-written (concise, impersonal, no first person); python-pptx is used
only to lay it out on the standard Microsoft Office template. The 11-model standings table
is pulled live from results/runs.jsonl via src/analysis/aggregate.py so the numbers stay in
sync with the report, and the six figures are embedded from disk.

    python3 scripts/make_slides.py   ->   report/SLM_vs_LLM_Mid_Sem_Review.pptx
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

from src.analysis.aggregate import load_results_df, per_model  # noqa: E402
from src.config import load_config, resolve_path  # noqa: E402

# Standard Office theme accents.
ACCENT = RGBColor(0x1F, 0x38, 0x64)   # dark blue — titles / table header
ACCENT2 = RGBColor(0x2E, 0x74, 0xB5)  # medium blue — rules / lead-ins
INK = RGBColor(0x22, 0x22, 0x22)
MUTE = RGBColor(0x59, 0x59, 0x59)
BAND = RGBColor(0xEE, 0xF2, 0xF8)     # light row banding

# 16:9 widescreen — the current PowerPoint standard size.
SW = Inches(13.333)
SH = Inches(7.5)

# Fallback display names (canonical source is config.yaml `display:` via per_model).
DISPLAY = {
    "smollm2-360m": "SmolLM2-360M", "qwen2.5-0.5b": "Qwen2.5-0.5B",
    "llama3.2-1b": "Llama-3.2-1B", "gemma3-1b": "Gemma-3-1B", "gemma2-2b": "Gemma-2-2B",
    "gemma3-4b": "Gemma-3-4B", "phi4-mini": "Phi-4-mini", "mistral-7b": "Mistral-7B",
    "gemma4-31b": "Gemma-4-31B", "large-open-llm": "Llama-3.3-70B",
    "frontier-llm": "GPT-OSS-120B",
}


# --------------------------------------------------------------------------- primitives
def _blank(prs):
    """A blank slide (layout 6) — full manual control, standard theme."""
    return prs.slides.add_slide(prs.slide_layouts[6])


def _textbox(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    return tb, tf


def _run(p, text, *, size=18, bold=False, italic=False, color=INK, font="Calibri"):
    r = p.add_run()
    r.text = text
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = font
    f.color.rgb = color
    return r


def _title_bar(slide, title, kicker=None):
    """Standard slide header: kicker (section), title, accent rule."""
    top = Inches(0.45)
    if kicker:
        _, ktf = _textbox(slide, Inches(0.6), Inches(0.28), Inches(12.1), Inches(0.35))
        p = ktf.paragraphs[0]
        _run(p, kicker.upper(), size=13, bold=True, color=ACCENT2)
        top = Inches(0.62)
    _, ttf = _textbox(slide, Inches(0.6), top, Inches(12.1), Inches(0.95))
    p = ttf.paragraphs[0]
    _run(p, title, size=30, bold=True, color=ACCENT)
    # accent rule
    line = slide.shapes.add_shape(1, Inches(0.62), Inches(1.62), Inches(2.2), Pt(3))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT2
    line.line.fill.background()
    line.shadow.inherit = False


def _bullets(slide, items, *, left=Inches(0.7), top=Inches(1.95),
             width=Inches(12.0), height=Inches(5.0), size=18, gap=6):
    """items: list of (text, level, lead) — lead is an optional bold blue lead-in phrase."""
    _, tf = _textbox(slide, left, top, width, height)
    tf.word_wrap = True
    first = True
    for it in items:
        text, level, lead = (it + (None,))[:3] if isinstance(it, tuple) else (it, 0, None)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(gap)
        p.space_before = Pt(0)
        bullet = "•   " if level == 0 else "–   "
        _run(p, bullet, size=size, color=ACCENT2, bold=True)
        if lead:
            _run(p, lead, size=size, bold=True, color=ACCENT)
            _run(p, text, size=size, color=INK)
        else:
            _run(p, text, size=size - (0 if level == 0 else 1),
                 color=INK if level == 0 else MUTE)
    return tf


def _caption(slide, text, *, top, left=Inches(0.6), width=Inches(12.1), size=13):
    _, tf = _textbox(slide, left, top, width, Inches(0.6))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, text, size=size, italic=True, color=MUTE)


def _image_fit(slide, path, *, left, top, max_w, max_h):
    """Embed keeping aspect, fit within max_w×max_h, centered in that box."""
    from PIL import Image
    try:
        iw, ih = Image.open(path).size
    except Exception:
        iw, ih = 1000, 620
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    off_l = left + (max_w - w) // 2
    off_t = top + (max_h - h) // 2
    slide.shapes.add_picture(str(path), Emu(off_l), Emu(off_t), Emu(w), Emu(h))


def _table(slide, headers, rows, *, left, top, width, col_w=None,
           header_size=11, body_size=11, first_left=True):
    nrows, ncols = len(rows) + 1, len(headers)
    height = Inches(0.34) * nrows
    gf = slide.shapes.add_table(nrows, ncols, left, top, width, height)
    tbl = gf.table
    if col_w:
        total = sum(col_w)
        for i, w in enumerate(col_w):
            tbl.columns[i].width = Emu(int(width * w / total))
    # header
    for c, h in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_top = cell.margin_bottom = Pt(2)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if (c == 0 and first_left) else PP_ALIGN.CENTER
        _run(p, h, size=header_size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    # body
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BAND if r % 2 else RGBColor(0xFF, 0xFF, 0xFF)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_top = cell.margin_bottom = Pt(1)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if (c == 0 and first_left) else PP_ALIGN.CENTER
            _run(p, str(val), size=body_size, bold=(c == 0 and first_left), color=INK)
    return tbl


# --------------------------------------------------------------------------- slides
def title_slide(prs, rep):
    slide = _blank(prs)
    # top accent band
    band = slide.shapes.add_shape(1, 0, 0, SW, Inches(0.28))
    band.fill.solid(); band.fill.fore_color.rgb = ACCENT2
    band.line.fill.background(); band.shadow.inherit = False

    _, tf = _textbox(slide, Inches(0.9), Inches(1.7), Inches(11.5), Inches(2.2))
    p = tf.paragraphs[0]
    _run(p, "SLM vs LLM for Key Information Extraction", size=40, bold=True, color=ACCENT)
    p2 = tf.add_paragraph()
    _run(p2, "from Financial Documents", size=40, bold=True, color=ACCENT)
    p3 = tf.add_paragraph(); p3.space_before = Pt(10)
    _run(p3, "Mid-Semester Review", size=24, bold=True, color=ACCENT2)

    rule = slide.shapes.add_shape(1, Inches(0.95), Inches(4.25), Inches(3.0), Pt(3))
    rule.fill.solid(); rule.fill.fore_color.rgb = ACCENT2
    rule.line.fill.background(); rule.shadow.inherit = False

    _, mf = _textbox(slide, Inches(0.9), Inches(4.5), Inches(11.5), Inches(2.3))
    lines = [
        (f"{rep['student_name']}   ·   BITS ID {rep['student_id']}", 20, True, INK),
        (f"{rep['degree_programme']}   ·   Course {rep['course_code']}", 17, False, MUTE),
        (f"Dissertation (BITS ZG628T)   ·   Supervisor: {rep['supervisor_name']}", 15, False, MUTE),
        ("Birla Institute of Technology & Science, Pilani   ·   July 2026", 15, False, MUTE),
    ]
    for i, (t, s, b, col) in enumerate(lines):
        p = mf.paragraphs[0] if i == 0 else mf.add_paragraph()
        p.space_after = Pt(6)
        _run(p, t, size=s, bold=b, color=col)


def agenda_slide(prs):
    slide = _blank(prs)
    _title_bar(slide, "Agenda")
    items = [
        ("Introduction — problem, gap, research question, objectives, scope", 0, None),
        ("Literature Review — four papers grounding the study", 0, None),
        ("Research Methodology — the text-in extraction pipeline and framework", 0, None),
        ("Data Collection — datasets, samples, test set, and the model registry", 0, None),
        ("Results to Date — the completed 11-model baseline (Phase 2)", 0, None),
        ("Key Findings & Future Plan — what Phase 3 will address", 0, None),
    ]
    _bullets(slide, items, top=Inches(2.1), size=21, gap=12)


def intro_slide(prs):
    slide = _blank(prs)
    _title_bar(slide, "Problem & Motivation", kicker="Introduction")
    items = [
        ("Large language models lead document extraction.", 0,
         "LLMs set the standard — "),
        ("but they are costly per call, high-latency, and cloud-only.", 1, None),
        ("Banks and firms often cannot send data out.", 0, "Data cannot leave the building — "),
        ("confidentiality and compliance frequently prohibit sending customer documents "
         "to an external API.", 1, None),
        ("Accurate extraction that runs locally, cheaply, and privately on an "
         "organisation's own hardware.", 0, "The real need — "),
    ]
    _bullets(slide, items, top=Inches(2.1), size=20, gap=14)


def gap_slide(prs):
    slide = _blank(prs)
    _title_bar(slide, "Research Gap", kicker="Introduction")
    items = [
        ("Recent work argues that small models — often a few billion parameters — are "
         "“good enough” for narrow, repetitive tasks, and far cheaper to run.", 0,
         "The claim:  "),
        ("The claim is unproven for real financial-document extraction.", 0, "The gap:  "),
        ("It is unknown how small a model can be before extraction breaks down.", 1, None),
        ("This dissertation tests the claim empirically, end to end.", 0, None),
    ]
    _bullets(slide, items, top=Inches(2.1), size=20, gap=16)


def rq_slide(prs):
    slide = _blank(prs)
    _title_bar(slide, "Research Question & Hypothesis", kicker="Introduction")
    items = [
        ("Can small, locally-run language models match large LLMs on financial key "
         "information extraction — at lower cost and latency, and with better data privacy?",
         0, "Research question:  "),
        ("SLMs match LLMs on simple documents.", 0, "Hypothesis:  "),
        ("The gap widens as documents grow more complex.", 1, None),
        ("SLMs win decisively on cost, latency, and privacy.", 1, None),
    ]
    _bullets(slide, items, top=Inches(2.1), size=20, gap=16)


def objectives_slide(prs):
    slide = _blank(prs)
    _title_bar(slide, "Objectives", kicker="Introduction")
    _, tf = _textbox(slide, Inches(0.7), Inches(1.95), Inches(12.0), Inches(0.8))
    p = tf.paragraphs[0]
    _run(p, "Determine whether small, locally-run models can match large LLMs for "
            "extracting key information from financial documents, and map the "
            "accuracy–efficiency trade-off.", size=17, italic=True, color=MUTE)
    items = [
        ("how closely small models approach large LLMs in extraction quality (F1).", 0,
         "Capability gap — "),
        ("how accuracy scales against cost, latency, and memory across model sizes.", 0,
         "Trade-off — "),
        ("the complexity point where small models stop being reliable.", 0,
         "Viability limit — "),
        ("the cost and privacy advantages of running models locally.", 0,
         "Enterprise value — "),
    ]
    _bullets(slide, items, top=Inches(2.85), height=Inches(4.4), size=19, gap=14)


def scope_slide(prs):
    slide = _blank(prs)
    _title_bar(slide, "Scope of Work", kicker="Introduction")
    items = [
        ("text-in key information extraction from financial documents "
         "(PDF → text → JSON).", 0, "Task — "),
        ("public datasets SROIE (simple) and Kleister-Charity (complex), both with "
         "gold labels.", 0, "Data — "),
        ("eleven models — eight small, locally-run SLMs and three large open-weight "
         "reference LLMs on serverless inference.", 0, "Models — "),
        ("zero-shot and few-shot prompting with identical settings for a fair comparison.",
         0, "Conditions — "),
        ("extraction quality (precision, recall, F1) and efficiency (latency, memory, cost).",
         0, "Evaluation — "),
    ]
    _bullets(slide, items, top=Inches(2.1), size=19, gap=13)


def litreview_overview_slide(prs):
    slide = _blank(prs)
    _title_bar(slide, "Literature Review — Four Papers", kicker="Literature Review")
    items = [
        ("Small Language Models are the Future of Agentic AI (Belcak et al., NVIDIA, 2025) "
         "— motivates using SLMs for narrow tasks.", 0, "Paper 1  ·  "),
        ("Are Small Language Models Ready to Compete with LLMs? (Sinha, Jain, Chadha, 2024) "
         "— a template for structured small-vs-large evaluation.", 0, "Paper 2  ·  "),
        ("Kleister: KIE Datasets with Long, Complex Documents (Stanisławek et al., "
         "ICDAR 2021) — the complex financial benchmark.", 0, "Paper 3  ·  "),
        ("SROIE: Scanned Receipt OCR and Information Extraction (Huang et al., ICDAR 2019) "
         "— the simple receipt benchmark and scoring protocol.", 0, "Paper 4  ·  "),
    ]
    _bullets(slide, items, top=Inches(2.1), size=19, gap=16)


def paper_slide(prs, n, title, cite, problem, did, finding, matters):
    slide = _blank(prs)
    _title_bar(slide, title, kicker=f"Literature Review · Paper {n}")
    _, cf = _textbox(slide, Inches(0.7), Inches(1.72), Inches(12.0), Inches(0.5))
    _run(cf.paragraphs[0], cite, size=14, italic=True, color=ACCENT2)
    items = [
        (problem, 0, "The problem:  "),
        (did, 0, "What they did:  "),
        (finding, 0, "Key finding:  "),
        (matters, 0, "Why it matters here:  "),
    ]
    _bullets(slide, items, top=Inches(2.35), size=18, gap=14)


def methodology_slide(prs, fig_dir):
    slide = _blank(prs)
    _title_bar(slide, "Text-in Extraction Pipeline", kicker="Research Methodology")
    fig = fig_dir / "fig2_pipeline.png"
    if fig.exists():
        _image_fit(slide, fig, left=Inches(0.8).emu, top=Inches(2.0).emu,
                   max_w=Inches(11.7).emu, max_h=Inches(3.6).emu)
    _caption(slide, "A single document flows: extract text → build prompt (fields + JSON schema) "
                    "→ run model → parse JSON → compare to gold.", top=Inches(5.75))
    _, tf = _textbox(slide, Inches(0.9), Inches(6.35), Inches(11.5), Inches(0.7))
    _run(tf.paragraphs[0],
         "Text only — no image or vision input. The same prompt and settings go to every "
         "model, which isolates the language model's extraction ability from OCR quality.",
         size=15, color=MUTE)


def architecture_slide(prs, fig_dir):
    slide = _blank(prs)
    _title_bar(slide, "Evaluation Framework Architecture", kicker="Research Methodology")
    fig = fig_dir / "fig1_architecture.png"
    if fig.exists():
        _image_fit(slide, fig, left=Inches(0.6).emu, top=Inches(1.95).emu,
                   max_w=Inches(7.4).emu, max_h=Inches(5.0).emu)
    items = [
        ("dataset loaders + gold handling", 0, "Data layer — "),
        ("normalise OCR output into clean lines", 0, "Preprocess — "),
        ("instruction + schema + document text", 0, "Prompt builder — "),
        ("Ollama (local) and DigitalOcean (cloud)", 0, "Registry + runners — "),
        ("read JSON, check against schema", 0, "Parser — "),
        ("precision, recall, F1, exact match", 0, "Metrics — "),
        ("latency, peak memory, throughput", 0, "Efficiency — "),
        ("aggregation, plots, summary tables", 0, "Analysis — "),
    ]
    _bullets(slide, items, left=Inches(8.2), top=Inches(2.0),
             width=Inches(4.9), size=15, gap=8)


def datasets_slide(prs):
    slide = _blank(prs)
    _title_bar(slide, "Datasets — Simple to Complex", kicker="Data Collection")
    left = [
        ("scanned receipts · simple baseline", 0, "SROIE — "),
        ("4 fields: company, date, address, total", 1, None),
        ("short, single-page documents", 1, None),
        ("scored with entity-level F1", 1, None),
    ]
    right = [
        ("UK charity annual reports · complex", 0, "Kleister-Charity — "),
        ("8 fields incl. monetary amounts & dates", 1, None),
        ("long, multi-page documents", 1, None),
        ("realistic, bank-like extraction", 1, None),
    ]
    _bullets(slide, left, left=Inches(0.7), top=Inches(2.1), width=Inches(6.0),
             size=18, gap=10)
    _bullets(slide, right, left=Inches(6.9), top=Inches(2.1), width=Inches(6.0),
             size=18, gap=10)
    _, tf = _textbox(slide, Inches(0.7), Inches(5.9), Inches(12.0), Inches(0.8))
    p = tf.paragraphs[0]
    _run(p, "The simple → complex pairing reveals exactly where the SLM–LLM gap opens up.",
         size=17, bold=True, color=ACCENT2)


def sample_slide(prs):
    slide = _blank(prs)
    _title_bar(slide, "Dataset Sample & Test Set", kicker="Data Collection")
    # SROIE example
    _, lf = _textbox(slide, Inches(0.7), Inches(1.95), Inches(6.0), Inches(3.6))
    _run(lf.paragraphs[0], "SROIE — input → gold output", size=16, bold=True, color=ACCENT)
    for t in ["BOOK TA .K (TAMAN DAYA) SDN BHD", "DATE: 25/12/2018   TOTAL: 9.00", "",
              '{ "company": "BOOK TA .K ...",', '  "date": "25/12/2018",',
              '  "address": "NO.53 ... JOHOR BAHRU",', '  "total": "9.00" }']:
        p = lf.add_paragraph(); p.space_after = Pt(2)
        _run(p, t, size=13, color=INK, font="Consolas")
    # Kleister example
    _, rf = _textbox(slide, Inches(6.9), Inches(1.95), Inches(6.0), Inches(3.6))
    _run(rf.paragraphs[0], "Kleister-Charity — gold fields (8)", size=16, bold=True, color=ACCENT)
    for t in ['{ "charity_name": "Smith Family Foundation",', '  "charity_number": "1093345",',
              '  "report_date": "2017-03-31",', '  "income_annually...": "1245000.00",',
              '  "spending_annually...": "1180500.00",', '  "address__post_town": "LONDON",',
              '  "address__postcode": "SW1A 1AA" }']:
        p = rf.add_paragraph(); p.space_after = Pt(2)
        _run(p, t, size=13, color=INK, font="Consolas")
    items = [
        ("20 documents, seed-fixed (10 SROIE + 10 Kleister), frozen across the study.", 0,
         "Test set — "),
        ("zero-shot / few-shot (k=2) × lines-only / lines+key-values = 80 measurements "
         "per model, 880 in total.", 0, "Condition grid — "),
    ]
    _bullets(slide, items, top=Inches(5.6), height=Inches(1.7), size=16, gap=8)


def models_slide(prs):
    slide = _blank(prs)
    _title_bar(slide, "The Models — Specs & Access", kicker="Data Collection")
    headers = ["Model", "Vendor", "Params (B)", "License", "Arm"]
    rows = [
        ["SmolLM2-360M", "Hugging Face", "0.36", "Apache-2.0", "Local"],
        ["Qwen2.5-0.5B", "Alibaba", "0.5", "Apache-2.0", "Local"],
        ["Gemma-3-1B", "Google", "1", "Gemma", "Local"],
        ["Llama-3.2-1B", "Meta", "1.24", "Llama 3.2", "Local"],
        ["Gemma-2-2B", "Google", "2", "Gemma", "Local"],
        ["Phi-4-mini", "Microsoft", "3.8", "MIT", "Local"],
        ["Gemma-3-4B", "Google", "4", "Gemma", "Local"],
        ["Mistral-7B", "Mistral AI", "7", "Apache-2.0", "Local"],
        ["Gemma-4-31B", "Google", "31", "Gemma", "Cloud"],
        ["Llama-3.3-70B", "Meta", "70", "Llama 3.3", "Cloud"],
        ["GPT-OSS-120B", "OpenAI", "120", "Apache-2.0", "Cloud"],
    ]
    _table(slide, headers, rows, left=Inches(2.1), top=Inches(1.95), width=Inches(9.1),
           col_w=[2.4, 2.0, 1.3, 1.8, 1.0], header_size=13, body_size=12)
    _, tf = _textbox(slide, Inches(0.7), Inches(6.7), Inches(12.0), Inches(0.5))
    _run(tf.paragraphs[0], "Eight local SLMs run through Ollama on an Apple M1 (8 GB); three "
         "large open-weight models run on DigitalOcean serverless inference.",
         size=14, italic=True, color=MUTE)


def metrics_slide(prs):
    slide = _blank(prs)
    _title_bar(slide, "Metrics — Quality & Efficiency", kicker="Research Methodology")
    left = [
        ("of returned fields, how many are correct.", 0, "Precision — "),
        ("of the true fields, how many are found.", 0, "Recall — "),
        ("balanced P/R — the KIE standard.", 0, "F1 — "),
        ("every field of a document correct.", 0, "Exact match — "),
    ]
    right = [
        ("time taken per document.", 0, "Latency — "),
        ("peak RAM/VRAM while running.", 0, "Memory — "),
        ("money per 1,000 documents processed.", 0, "Cost — "),
        ("share of unparseable outputs.", 0, "Parse-fail rate — "),
    ]
    _, h1 = _textbox(slide, Inches(0.7), Inches(1.85), Inches(6.0), Inches(0.5))
    _run(h1.paragraphs[0], "Extraction quality", size=18, bold=True, color=ACCENT2)
    _, h2 = _textbox(slide, Inches(6.9), Inches(1.85), Inches(6.0), Inches(0.5))
    _run(h2.paragraphs[0], "Operational efficiency", size=18, bold=True, color=ACCENT2)
    _bullets(slide, left, left=Inches(0.7), top=Inches(2.5), width=Inches(6.0),
             size=18, gap=12)
    _bullets(slide, right, left=Inches(6.9), top=Inches(2.5), width=Inches(6.0),
             size=18, gap=12)


def standings_slide(prs, pm):
    slide = _blank(prs)
    _title_bar(slide, "Results — 11-Model Standings", kicker="Results to Date · Phase 2")
    pms = pm.sort_values("f1_macro", ascending=False)
    rows = []
    for _, r in pms.iterrows():
        is_local = r["model_type"] == "local"
        cost = (f"${r['cost_low_per_doc']:.5f}–{r['cost_high_per_doc']:.5f}*"
                if is_local else f"${r['cost_usd_per_doc']:.5f}")
        rows.append([
            r.get("display_name") or DISPLAY.get(r["model_id"], r["model_id"]),
            f"{r['params_b']:g}",
            "local" if is_local else "cloud",
            f"{r['f1_zero_shot']:.3f}" if r["f1_zero_shot"] is not None else "—",
            f"{r['f1_few_shot']:.3f}" if r["f1_few_shot"] is not None else "—",
            f"{r['precision']:.2f}", f"{r['recall']:.2f}",
            f"{r['exact_match_rate']:.3f}", f"{r['latency_s']:.1f}", cost,
        ])
    headers = ["Model", "Params", "Arm", "F1 zero", "F1 few", "P", "R", "Exact",
               "Lat (s)", "Cost/doc"]
    _table(slide, headers, rows, left=Inches(0.5), top=Inches(1.85), width=Inches(12.3),
           col_w=[2.0, 0.9, 0.8, 0.9, 0.9, 0.6, 0.6, 0.9, 0.9, 2.0],
           header_size=11, body_size=11)
    _, tf = _textbox(slide, Inches(0.5), Inches(6.55), Inches(12.3), Inches(0.7))
    p = tf.paragraphs[0]
    _run(p, "Quality does not track parameter count. ", size=14, bold=True, color=ACCENT)
    _run(p, "Ordered by few-shot F1. *Local cost is a market band (rate to rent equivalent "
            "compute), not an on-device charge.", size=13, color=MUTE)


def two_figure_slide(prs, title, kicker, fig_a, fig_b, cap_a, cap_b, note):
    slide = _blank(prs)
    _title_bar(slide, title, kicker=kicker)
    for fig, lft, cap in [(fig_a, Inches(0.55), cap_a), (fig_b, Inches(6.9), cap_b)]:
        if fig and fig.exists():
            _image_fit(slide, fig, left=lft.emu, top=Inches(1.95).emu,
                       max_w=Inches(6.0).emu, max_h=Inches(3.9).emu)
        _caption(slide, cap, top=Inches(5.85), left=lft, width=Inches(6.0), size=12)
    _, tf = _textbox(slide, Inches(0.7), Inches(6.5), Inches(12.0), Inches(0.8))
    _run(tf.paragraphs[0], note, size=15, color=INK)


def heatmap_slide(prs, plots_dir):
    slide = _blank(prs)
    _title_bar(slide, "Results — Effects & Per-field Accuracy", kicker="Results to Date · Phase 2")
    fig = plots_dir / "plot_field_heatmap.png"
    if fig.exists():
        _image_fit(slide, fig, left=Inches(0.55).emu, top=Inches(1.95).emu,
                   max_w=Inches(7.2).emu, max_h=Inches(4.7).emu)
    items = [
        ("removed every output-format failure of the smallest model (31 of 80 → 0) and "
         "lifted its F1 in step.", 0, "Constrained decoding — "),
        ("helps the larger models (Llama-3.3 +0.13) but can collapse the smallest "
         "(Llama-3.2-1B 0.53 → 0.03, copying the example).", 0, "Few-shot — "),
        ("simple receipt fields extract reliably; missed recall concentrates in the hard "
         "Kleister income and spending figures.", 0, "Per-field — "),
    ]
    _bullets(slide, items, left=Inches(8.0), top=Inches(2.05), width=Inches(5.1),
             size=15, gap=12)
    _caption(slide, "Figure 6: Per-field accuracy heatmap.", top=Inches(6.75),
             left=Inches(0.55), width=Inches(7.2), size=12)


def findings_slide(prs):
    slide = _blank(prs)
    _title_bar(slide, "Key Findings & Future Plan", kicker="Summary")
    items = [
        ("a 4B local model matches the 120B cloud model, and a 7B local model overtakes it.",
         0, "Parameter count ≠ quality: "),
        ("it lifts small models the most by removing malformed output.", 0,
         "Constrained decoding is the equaliser: "),
        ("every competitive local model dominates a paid model of similar F1 on cost, "
         "with no data leaving the machine.", 0, "Local-first wins on economics: "),
    ]
    _bullets(slide, items, top=Inches(1.95), width=Inches(12.2), size=17, gap=10)
    headers = ["Phase", "Window", "Work", "Status"]
    rows = [
        ["Baseline (Phase 2)", "Feb–Jun 2026", "Framework + 11-model baseline", "COMPLETED"],
        ["Fine-tuning (Phase 3)", "Jul–Sep 2026", "QLoRA on phi4-mini + mistral-7b", "PENDING"],
        ["Recall lift (Phase 3)", "Sep–Oct 2026", "Self-consistency + regex recovery", "PENDING"],
        ["Write-up & viva", "Oct–Nov 2026", "Final report and defence", "PENDING"],
    ]
    _table(slide, headers, rows, left=Inches(0.7), top=Inches(4.35), width=Inches(11.9),
           col_w=[2.4, 1.7, 4.6, 1.4], header_size=12, body_size=12)


def thanks_slide(prs, rep):
    slide = _blank(prs)
    band = slide.shapes.add_shape(1, 0, Inches(3.0), SW, Inches(1.5))
    band.fill.solid(); band.fill.fore_color.rgb = ACCENT
    band.line.fill.background(); band.shadow.inherit = False
    tf = band.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, "Thank You  ·  Questions & Discussion", size=32, bold=True,
         color=RGBColor(0xFF, 0xFF, 0xFF))
    _, sf = _textbox(slide, Inches(0.9), Inches(4.7), Inches(11.5), Inches(0.8))
    q = sf.paragraphs[0]; q.alignment = PP_ALIGN.CENTER
    _run(q, f"{rep['student_name']}   ·   {rep['student_id']}   ·   Mid-Semester Review, July 2026",
         size=17, color=MUTE)


# --------------------------------------------------------------------------- assembly
def main() -> None:
    cfg = load_config()
    rep = cfg["report"]
    plots_dir = resolve_path(cfg, cfg["paths"]["plots_dir"])
    report_dir = resolve_path(cfg, cfg["paths"]["report_dir"])
    fig_dir = report_dir / "figures"

    df = load_results_df(cfg)
    pm = per_model(df[df["shot_mode"] == "few_shot"], cfg)
    zero_f1 = per_model(df, cfg).set_index("model_id")["f1_zero_shot"]
    pm = pm.copy()
    pm["f1_zero_shot"] = pm["model_id"].map(zero_f1)

    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH

    title_slide(prs, rep)
    agenda_slide(prs)
    intro_slide(prs)
    gap_slide(prs)
    rq_slide(prs)
    objectives_slide(prs)
    scope_slide(prs)
    litreview_overview_slide(prs)
    paper_slide(prs, 1, "Small Language Models are the Future of Agentic AI",
                "Belcak, Heinrich, Diao, Fu, Dong, Muralidharan, Lin, Molchanov (NVIDIA) "
                "— arXiv:2506.02153, 2025",
                "LLMs are the default, yet most real applications run a few narrow, repetitive "
                "tasks — overkill for a giant general model.",
                "A position paper arguing from capability, agent architecture, and economics "
                "that SLMs are the better fit.",
                "SLMs are capable enough, far cheaper (often 10–30×), lower latency, and "
                "can run on edge or on-premise devices.",
                "Motivates testing small models for a narrow task (financial KIE) instead of "
                "defaulting to a large LLM.")
    paper_slide(prs, 2, "Are Small Language Models Ready to Compete with LLMs?",
                "Sinha, Jain, Chadha — arXiv:2406.11402, 2024 (rev. 2025)",
                "State-of-the-art LLMs are often infeasible (cost, size, proprietary), but small "
                "models do not perform well universally.",
                "Proposes a framework to evaluate small open models across task types, domains, "
                "and prompt styles; compares ten SLMs.",
                "Chosen appropriately, small models can beat some SOTA LLMs and even compete with "
                "GPT-4-class models.",
                "A direct methodological template for this structured, prompt-aware "
                "small-vs-large comparison.")
    paper_slide(prs, 3, "Kleister: KIE Datasets with Long, Complex-Layout Documents",
                "Stanisławek et al. — Proceedings of ICDAR 2021 — arXiv:2105.05796",
                "Real KIE involves long, multi-page documents with complex layouts — harder than "
                "short receipts or forms.",
                "Introduces the Kleister datasets (incl. Kleister-Charity from financial reports) "
                "with names, dates, and monetary values.",
                "Provides gold-annotated benchmarks and shows that long-document extraction is "
                "genuinely challenging.",
                "Kleister-Charity is the complex, financial, multi-page dataset and ground-truth "
                "source for this study.")
    paper_slide(prs, 4, "SROIE: Scanned Receipt OCR and Information Extraction",
                "Huang, Chen, He, Bai, Karatzas, Lu, Jawahar — ICDAR 2019 Competition "
                "— arXiv:1905.13538",
                "Extracting a few key fields from scanned receipts is a common but non-trivial "
                "task, and lacked a shared, labelled benchmark.",
                "Introduces the SROIE task and dataset — receipts with gold labels for company, "
                "date, address, and total — and an entity-level F1 scoring protocol.",
                "Establishes a widely-used benchmark; the constrained field set makes it a clean "
                "measure of extraction accuracy.",
                "SROIE is the simple dataset and the scoring protocol used here — the complexity "
                "counterpart to Kleister.")
    methodology_slide(prs, fig_dir)
    architecture_slide(prs, fig_dir)
    datasets_slide(prs)
    sample_slide(prs)
    models_slide(prs)
    metrics_slide(prs)
    standings_slide(prs, pm)
    two_figure_slide(
        prs, "Results — Quality vs Size & Cost", "Results to Date · Phase 2",
        plots_dir / "plot_f1_vs_params.png", plots_dir / "plot_cost_frontier.png",
        "Figure 3: F1 versus model size (log axis).",
        "Figure 4: Quality versus cost per document.",
        "F1 is flat past a few billion parameters; every competitive local model is several "
        "times cheaper than a cloud model of similar quality.")
    two_figure_slide(
        prs, "Results — Latency vs Size", "Results to Date · Phase 2",
        plots_dir / "plot_latency_vs_params.png", None,
        "Figure 5: Latency versus model size (local models).", "",
        "Small models answer in 3–6 s; the 7B model is the slowest local model because it "
        "pushes against the 8 GB memory limit and pages to disk.")
    heatmap_slide(prs, plots_dir)
    findings_slide(prs)
    thanks_slide(prs, rep)

    out = report_dir / "SLM_vs_LLM_Mid_Sem_Review.pptx"
    prs.save(str(out))
    print(f"Wrote {out}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
