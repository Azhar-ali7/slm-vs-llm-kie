"""Assemble the BITS mid-semester report (Word .docx) in the prescribed format.

The prose in this file is hand-written (not auto-generated findings); python-docx is
used only to lay it out in the BITS structure. Result tables are pulled live from
results/runs.jsonl via src/analysis/aggregate.py so the numbers stay in sync.

    python3 scripts/make_report_doc.py   ->   report/Mid_Sem_Report.docx
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

from src.analysis.aggregate import load_results_df, per_model  # noqa: E402
from src.config import load_config, resolve_path  # noqa: E402

# Fallback display names. The canonical source is the `display:` field on each
# model in config.yaml (carried through registry.model_meta -> per_model's
# display_name column); this dict is only used if that is missing.
DISPLAY = {
    "smollm2-360m": "SmolLM2-360M",
    "qwen2.5-0.5b": "Qwen2.5-0.5B",
    "llama3.2-1b": "Llama-3.2-1B",
    "gemma3-1b": "Gemma-3-1B",
    "gemma2-2b": "Gemma-2-2B",
    "gemma3-4b": "Gemma-3-4B",
    "phi4-mini": "Phi-4-mini",
    "mistral-7b": "Mistral-7B",
    "gemma4-31b": "Gemma-4-31B",
    "large-open-llm": "Llama-3.3-70B",
    "frontier-llm": "GPT-OSS-120B",
}


# --------------------------------------------------------------------------- helpers
def _set_base_style(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15


def _center(p) -> None:
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _para(doc, text="", *, bold=False, italic=False, size=None, center=False, space_after=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if center:
        _center(p)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def _heading(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14 if level == 1 else 12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p


def _figure(doc, path: Path, number: int, caption: str, width=5.8):
    if path.exists():
        p = doc.add_paragraph()
        _center(p)
        p.add_run().add_picture(str(path), width=Inches(width))
    else:  # graceful fallback if a figure is missing
        _para(doc, f"[figure not found: {path.name}]", italic=True, center=True)
    cap = doc.add_paragraph()
    _center(cap)
    r = cap.add_run(f"Figure {number}: {caption}")
    r.italic = True
    r.font.size = Pt(10)


def _table(doc, headers, rows, *, number=None, caption=None, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(str(h))
        run.bold = True
        run.font.size = Pt(10)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run("" if val is None else str(val))
            run.font.size = Pt(10)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    if number is not None and caption:
        cap = doc.add_paragraph()
        _center(cap)
        r = cap.add_run(f"Table {number}: {caption}")
        r.italic = True
        r.font.size = Pt(10)
    return table


# --------------------------------------------------------------------------- sections
def title_page(doc, rep):
    _para(doc, "", space_after=24)
    _para(doc, rep["title"].upper(), bold=True, size=18, center=True, space_after=18)
    _para(doc, "BITS ZG628T: Dissertation", size=12, center=True, space_after=18)
    _para(doc, "by", size=12, center=True, space_after=18)
    _para(doc, rep["student_name"], size=14, center=True, space_after=2)
    _para(doc, rep["student_id"], size=11, center=True, space_after=24)
    _para(doc, "Dissertation work carried out at", size=12, center=True, space_after=10)
    _para(doc, rep["organization"], size=12, center=True, space_after=24)
    _para(doc, f"Submitted in partial fulfilment of {rep['degree_programme']}", size=12,
          center=True, space_after=2)
    _para(doc, "degree programme", size=12, center=True, space_after=24)
    _para(doc, "Under the Supervision of", size=12, center=True, space_after=8)
    _para(doc, rep["supervisor_name"], size=12, center=True, space_after=4)
    _para(doc, rep["organization"], size=12, center=True, space_after=24)
    _para(doc, "(BITS Pilani logo — insert institute logo here)", italic=True, size=9,
          center=True, space_after=24)
    _para(doc, "BIRLA INSTITUTE OF TECHNOLOGY & SCIENCE", bold=True, size=12, center=True,
          space_after=4)
    _para(doc, "PILANI (RAJASTHAN)", size=12, center=True, space_after=4)
    _para(doc, rep["month_year"], size=12, center=True)
    doc.add_page_break()


def abstract(doc, rep):
    _heading(doc, "ABSTRACT")
    _para(doc,
          "Key Information Extraction (KIE) is the task of turning an unstructured "
          "document — a receipt, an invoice, or an annual report — into a small set of "
          "structured fields such as the company name, the date, the total amount, or the "
          "income reported for a year. Large language models handle this task well, but "
          "running them as a cloud service raises three practical concerns: the per-document "
          "cost, the latency of a network round trip, and the need to send commercially "
          "sensitive financial data to a third party. This dissertation studies whether small "
          "language models that run locally on ordinary hardware can come close to the "
          "accuracy of large cloud models on KIE from financial documents.")
    _para(doc,
          "A reproducible evaluation framework was built to compare eleven models, ranging "
          "from 0.36 billion to 120 billion parameters, on two public datasets: SROIE "
          "(scanned receipts, four fields) and Kleister-Charity (United Kingdom charity "
          "annual reports, eight fields). The eight small and mid-sized models run locally "
          "through Ollama on an Apple M1 machine with 8 GB of memory; the three large models "
          "run on DigitalOcean serverless inference. Each model is evaluated on the same "
          "20-document test set under four prompting conditions — zero-shot and few-shot, "
          "each combined with two input representations — which gives 80 measurements per "
          "model. Output quality is measured with precision, recall, F1, and document-level "
          "exact match; efficiency is measured with latency, peak memory, and cost.")
    _para(doc,
          "The measurements show that accuracy does not follow parameter count. Averaged across "
          "the four conditions, the strongest model is Gemma-4-31B, with an F1 of 0.70, which is "
          "higher than both the 70-billion and the 120-billion parameter models. Among the local "
          "models, a 4-billion parameter model (Gemma-3-4B, F1 0.64) matches the 120-billion "
          "parameter cloud model (F1 0.66) at a fraction of the cost and without any data leaving "
          "the machine. "
          "Constrained JSON decoding, applied to the local models, removed every "
          "output-format failure of the smallest model — from 31 failures out of 80 down to "
          "none — and raised its F1 in step. Recall is lower than precision for every model, "
          "which shows that the main source of error is missed fields rather than wrong ones.")
    _para(doc,
          "This report covers the mid-semester phase of the work: the framework, the two "
          "datasets, and the complete baseline comparison across all eleven models. "
          "Task-specific fine-tuning of the small models, which is expected to close the "
          "remaining gap to the cloud models, is reserved for the final phase and is set out "
          "in the future plan in Section 5.")

    _para(doc, "", space_after=18)
    # Signature block (two columns)
    sig = doc.add_table(rows=4, cols=2)
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = [("Signature of the Student", "Signature of the Supervisor"),
              (f"Name: {rep['student_name']}", "Name:"),
              ("Date:", "Date:"),
              ("Place:", "Place:")]
    for r, (left, right) in enumerate(labels):
        for c, text in ((0, left), (1, right)):
            cell = sig.rows[r].cells[c]
            cell.text = ""
            run = cell.paragraphs[0].add_run(text)
            run.font.size = Pt(10)
            if r == 0:
                run.bold = True
    doc.add_page_break()


def contents(doc):
    _heading(doc, "Contents")
    sections = [
        "1. Modules in the Evaluation Framework",
        "2. Functional Block Diagram / Description",
        "3. Major Technical Specifications and Results",
        "4. Design Considerations",
        "5. Future Plan",
        "6. Abbreviations",
    ]
    for s in sections:
        doc.add_paragraph(s, style="List Bullet" if False else None)
    _para(doc, "", space_after=6)
    _para(doc, "List of Figures", bold=True)
    figs = [
        "Figure 1: Modules of the evaluation framework",
        "Figure 2: Functional pipeline for a single document",
        "Figure 3: F1 versus model size",
        "Figure 4: Quality versus cost",
        "Figure 5: Latency versus model size",
        "Figure 6: Per-field accuracy heatmap",
    ]
    for f in figs:
        doc.add_paragraph(f)
    _para(doc, "", space_after=6)
    _para(doc, "List of Tables", bold=True)
    for t in ["Table 1: Experimental bench specifications",
              "Table 2: Baseline results for all eleven models",
              "Table 3: Future plan"]:
        doc.add_paragraph(t)
    doc.add_page_break()


def section_modules(doc, fig_dir):
    _heading(doc, "1. Modules in the Evaluation Framework")
    _para(doc,
          "The framework is organised as a set of independent modules, each responsible for "
          "one stage of the experiment. Keeping the stages separate makes it possible to swap "
          "a dataset, add a model, or change a prompt without touching the rest of the system. "
          "The modules and the flow of data between them are shown in Figure 1.")
    _para(doc, "The system consists of the following modules:")
    mods = [
        "(a) Data layer — dataset loaders and ground-truth handling",
        "(b) Preprocessing — normalising OCR output into clean text lines",
        "(c) Prompt builder — assembling the instruction, schema, and document text",
        "(d) Model registry and runners — local (Ollama) and cloud (DigitalOcean) backends",
        "(e) Output parser — reading the model's JSON and checking it against the schema",
        "(f) Metrics — precision, recall, F1, and exact match against the gold fields",
        "(g) Efficiency monitor — latency, peak memory, and token throughput",
        "(h) Analysis and reporting — aggregation, plots, and summary tables",
    ]
    for m in mods:
        doc.add_paragraph(m)
    _figure(doc, fig_dir / "fig1_architecture.png", 1,
            "Modules of the evaluation framework")
    _para(doc, "The function of each module is described below.")
    descs = [
        ("(a) Data layer", "Two datasets are supported through a common loader interface. SROIE "
         "provides scanned shop receipts with four target fields (company, date, address, and "
         "total). Kleister-Charity provides United Kingdom charity annual reports with eight "
         "fields, including the charity name, registration number, reporting period, and the "
         "income and spending figures. The loader reads each document together with its gold "
         "answer so that predictions can be scored automatically."),
        ("(b) Preprocessing", "Each document arrives already OCR'd, in one of two forms: "
         "positioned words with bounding boxes (SROIE receipts) or flat OCR text "
         "(Kleister-Charity reports). Word-based records are grouped into reading-order lines "
         "by clustering on vertical position; text-based records are split into lines and the "
         "financially relevant ones selected. Both produce the same simple JSON — a list of "
         "text lines and, for the receipt data, an optional set of key–value pairs. Long "
         "reports are capped at a fixed number of lines so that the prompt stays within the "
         "context window of the smallest models and the cost of the cloud models stays "
         "bounded."),
        ("(c) Prompt builder", "For each document the prompt builder produces an instruction, "
         "the target schema (the exact field names expected in the answer), and the document "
         "text. It supports two input representations — text lines only, and text lines plus "
         "extracted key–value pairs — and two prompting modes — zero-shot, and few-shot with "
         "two worked examples drawn from the training split. The examples are never taken from "
         "the test set."),
        ("(d) Model registry and runners", "The registry maps a model identifier to a runner. "
         "Local models run through Ollama on the M1 machine; large models run through "
         "DigitalOcean serverless inference over an OpenAI-compatible interface. The local "
         "runner additionally passes a JSON schema to the decoder so that the output is forced "
         "to be valid JSON with all required fields present. Temperature is fixed at zero and "
         "the random seed is fixed so that runs are repeatable."),
        ("(e) Output parser", "The parser reads the model's response, extracts the JSON object, "
         "and checks it against the schema. A response that cannot be parsed is recorded as a "
         "parse failure rather than silently dropped, so that output reliability can be "
         "measured as a metric in its own right."),
        ("(f) Metrics", "Each predicted field is compared with the gold value to give counts of "
         "correct, wrong, and missing fields. From these counts the framework computes "
         "precision, recall, and F1, both per field and aggregated, along with a document-level "
         "exact-match rate that is satisfied only when every field of a document is correct."),
        ("(g) Efficiency monitor", "While a model runs, the framework records the wall-clock "
         "latency of each call, the peak resident memory of the local server, the token "
         "throughput, and, for cloud models, the token cost. These figures support the "
         "quality-versus-cost analysis in Section 3."),
        ("(h) Analysis and reporting", "The analysis module aggregates the per-document results "
         "into per-model and per-condition tables, produces the plots in Section 3, and writes "
         "the summary files that this report draws on."),
    ]
    for name, text in descs:
        p = doc.add_paragraph()
        r = p.add_run(name + ": ")
        r.bold = True
        p.add_run(text)


def section_pipeline(doc, fig_dir):
    _heading(doc, "2. Functional Block Diagram / Description")
    _para(doc,
          "The framework processes one document at a time through a fixed sequence of steps. "
          "A document is first preprocessed into clean text lines. The prompt builder then "
          "combines these lines with the instruction and the target schema to form the prompt "
          "for the chosen condition. The prompt is sent to the model: for a local model the "
          "decoder is constrained to emit valid JSON, while for a cloud model the same "
          "structure is requested through the instruction. The model's reply is parsed into a "
          "JSON object, compared field by field against the gold answer, and scored. The scores "
          "and the efficiency measurements are written to a results file, one line per "
          "document. The pipeline is shown in Figure 2.")
    _figure(doc, fig_dir / "fig2_pipeline.png", 2,
            "Functional pipeline for a single document")
    _para(doc, "Two points about the experimental design are worth stating here.", )
    p = doc.add_paragraph()
    p.add_run("Experiment grid: ").bold = True
    p.add_run("every model is run over the same grid of 20 documents and four conditions "
              "(two prompting modes by two input representations), which produces 80 "
              "measurements per model and 880 measurements in total across the eleven models. "
              "Fixing the grid keeps the comparison fair: every model sees exactly the same "
              "inputs.")
    p = doc.add_paragraph()
    p.add_run("Constrained decoding and the baseline: ").bold = True
    p.add_run("the local models are evaluated with JSON-constrained decoding switched on. To "
              "measure the effect of that decision, an earlier run without the constraint is "
              "kept frozen as a baseline, and the two are compared directly. The constraint is "
              "a property of the local decoder, so it does not apply to the cloud models; that "
              "they still return clean JSON without it is itself a useful observation.")


def section_specs(doc, plots_dir, pm):
    _heading(doc, "3. Major Technical Specifications and Results")
    _para(doc,
          "The fixed specifications of the evaluation bench are listed in Table 1. These "
          "settings do not change between runs; they define the conditions under which every "
          "model is measured.")
    spec_rows = [
        ["1", "Local hardware", "Apple M1, 8 GB unified memory; data and code on external SSD"],
        ["2", "Local runtime", "Ollama, temperature 0.0, seed 42, output cap 512 tokens"],
        ["3", "Output decoding (local)", "JSON-schema constrained decoding (valid JSON enforced)"],
        ["4", "Large-model runtime", "DigitalOcean serverless inference (OpenAI-compatible)"],
        ["5", "Large models", "Gemma-4-31B, Llama-3.3-70B, GPT-OSS-120B (open weights)"],
        ["6", "Dataset 1", "SROIE — scanned receipts, 4 fields (simple)"],
        ["7", "Dataset 2", "Kleister-Charity — UK charity annual reports, 8 fields (complex)"],
        ["8", "Test set", "20 documents, seed-fixed (10 SROIE + 10 Kleister)"],
        ["9", "Conditions", "zero-shot / few-shot (k=2) × lines-only / lines+key-values"],
        ["10", "Measurements per model", "20 documents × 4 conditions = 80"],
        ["11", "Long-document cap", "120 input lines (bounds context and cost)"],
        ["12", "Quality metrics", "precision, recall, F1, document exact-match, parse-fail rate"],
        ["13", "Efficiency metrics", "latency, peak memory, token throughput, cost"],
    ]
    _table(doc, ["Sl. No", "Specification", "Value"], spec_rows,
           number=1, caption="Experimental bench specifications",
           widths=[0.7, 2.1, 4.0])

    _para(doc, "", space_after=6)
    _heading(doc, "3.2 Results to date", level=2)
    _para(doc,
          "Table 2 lists the baseline result for every model under two-shot prompting: F1 (few), "
          "precision, recall, exact match, latency, and cost are all measured on the same "
          "few-shot runs, so the row is internally consistent. The zero-shot F1 is carried in the "
          "first metric column as a reference, to show how each model responds to examples. The "
          "models are ordered by their few-shot F1. The clearest pattern is that quality does not "
          "rise with size. The top two scores are close — Llama-3.3-70B at 0.73 and Gemma-4-31B "
          "at 0.72 — yet Gemma-4-31B reaches its score with less than half the parameters and "
          "roughly a fifth of the latency (3.4 s against 16.3 s per document). More striking is "
          "that Mistral-7B, a seven-billion parameter model running locally on the M1, scores "
          "0.71 and so overtakes the 120-billion parameter GPT-OSS (0.67). Gemma-3-4B reaches "
          "0.63, close to that same 120-billion model, at a small fraction of the cloud "
          "per-document cost. A 2-billion parameter Gemma (0.65) is well ahead of the 3.8-billion "
          "parameter Phi-4-mini (0.47), which confirms that architecture and training matter more "
          "than raw parameter count at this scale.")
    _para(doc,
          "The split also exposes that few-shot prompting is not uniformly helpful. It lifts the "
          "larger models substantially — Llama-3.3 by 0.13 and Mistral-7B by 0.13 — but degrades "
          "several of the smallest ones. Llama-3.2-1B is the extreme case: a usable 0.53 zero-shot "
          "collapses to 0.03 with examples, because the model copies an example's answer instead "
          "of extracting from the target document. Reporting a single averaged F1 would hide this "
          "effect entirely, which is why the two columns are kept apart.")

    # Results table, sorted by F1 desc
    pms = pm.sort_values("f1_macro", ascending=False)
    res_rows = []
    for _, r in pms.iterrows():
        is_local = r["model_type"] == "local"
        # Local: hosted-rate band low–high (self-hosting has no per-token bill, but
        # "$0" misrepresents the economics — see Figure 4). Cloud: real billed cost.
        cost = (f"${r['cost_low_per_doc']:.5f}–{r['cost_high_per_doc']:.5f}*"
                if is_local else f"${r['cost_usd_per_doc']:.5f}")
        res_rows.append([
            r.get("display_name") or DISPLAY.get(r["model_id"], r["model_id"]),
            f"{r['params_b']:g}",
            "local" if is_local else "cloud",
            f"{r['f1_zero_shot']:.3f}" if r["f1_zero_shot"] is not None else "—",
            f"{r['f1_few_shot']:.3f}" if r["f1_few_shot"] is not None else "—",
            f"{r['precision']:.2f}",
            f"{r['recall']:.2f}",
            f"{r['exact_match_rate']:.3f}",
            f"{r['latency_s']:.1f}",
            cost,
        ])
    _table(doc,
           ["Model", "Params (B)", "Arm", "F1 (zero)", "F1 (few)", "P", "R", "Exact",
            "Latency (s)", "Cost/doc"],
           res_rows, number=2, caption="Baseline results for all eleven models",
           widths=[1.25, 0.6, 0.45, 0.55, 0.55, 0.35, 0.35, 0.5, 0.7, 1.25])
    _para(doc,
          "* Local models carry no per-token bill on the M1; the cost shown is a RANGE — the "
          "low-to-high commercial rate to host a model of that size (June 2026 provider rates) "
          "applied to the real token counts recorded for each run. It is the market price of the "
          "same compute, not an on-device charge.", space_after=6)

    _para(doc,
          "Figure 3 plots F1 against model size on a logarithmic axis. The curve is flat over a "
          "wide range: once a model reaches a few billion parameters, adding more brings little "
          "further gain on this task.")
    _figure(doc, plots_dir / "plot_f1_vs_params.png", 3, "F1 versus model size")
    _para(doc,
          "Figure 4 places quality against cost per document on a logarithmic axis. The billed "
          "DigitalOcean models are drawn at their real per-document cost; each local model is "
          "drawn as a cost RANGE (the horizontal bar) — the low-to-high per-token rate it would "
          "cost to rent a model of that size, since self-hosting has no per-token bill yet is not "
          "truly free. Even at the top of its band, every local model is several times cheaper "
          "than the cloud models of comparable F1, so a local model that reaches a competitive "
          "score dominates a paid model of similar quality.")
    _figure(doc, plots_dir / "plot_cost_frontier.png", 4, "Quality versus cost")
    _para(doc,
          "Figure 5 shows latency against model size for the local models. The smallest models "
          "answer in three to six seconds; the 7-billion parameter model is the slowest local "
          "model because it pushes against the 8 GB memory limit and pages to disk.")
    _figure(doc, plots_dir / "plot_latency_vs_params.png", 5, "Latency versus model size")
    _para(doc,
          "Figure 6 breaks accuracy down by field. The simple receipt fields are extracted "
          "reliably by most models, whereas the harder Kleister figures — reported income and "
          "spending — account for much of the lost recall and are the natural target for the "
          "next phase.")
    _figure(doc, plots_dir / "plot_field_heatmap.png", 6, "Per-field accuracy heatmap")


def section_design(doc):
    _heading(doc, "4. Design Considerations")
    points = [
        "The scope is text-in only. Documents are consumed as OCR text or line boxes; no image "
        "or vision model is used. This keeps the work within the memory budget of the M1 and "
        "isolates the language model's extraction ability from OCR quality.",
        "Constrained JSON decoding is used for the local models. It removes output-format "
        "failures at no extra cost and helps the smallest models the most, because their errors "
        "were largely malformed output rather than wrong content.",
        "The study is built to be reproducible. The seed, the 20-document test set, and the "
        "model settings are fixed, and each completed run is frozen as a versioned snapshot so "
        "that later changes can be measured as a before-and-after difference.",
        "Cost is controlled by running small and mid-sized models locally and using pay-per-token "
        "serverless inference for the large models. No GPU server is rented by the hour, and the "
        "entire large-model comparison cost under twenty cents.",
        "Long reports are bounded with a line cap and the local models are kept loaded between "
        "calls, so that the experiment runs within 8 GB of memory.",
        "The large arm uses open-weight models. This removes any dependency on closed models and "
        "makes the large-model results fully reproducible by anyone with the same setup.",
        "The training and test splits are kept strictly separate. This is needed so that the "
        "fine-tuning planned for the final phase can be measured honestly against the same "
        "untouched test set.",
    ]
    for pt in points:
        doc.add_paragraph(pt, style="List Bullet")


def section_future(doc):
    _heading(doc, "5. Future Plan")
    _para(doc,
          "The work is divided so that the mid-semester deliverable stands on its own as a "
          "complete baseline, while the improvement work is kept for the final phase. The plan "
          "is summarised in Table 3.")
    rows = [
        ["1", "Abstract and proposal", "Jan 2026 – Feb 2026",
         "Define the problem, datasets, and evaluation plan", "COMPLETED"],
        ["2", "Framework and baseline (Phase 2)", "Mar 2026 – Jun 2026",
         "Build the framework; run the 11-model baseline across both datasets", "COMPLETED"],
        ["3", "Fine-tuning (Phase 3)", "Jul 2026 – Aug 2026",
         "QLoRA fine-tuning of Phi-4-mini and Mistral-7B on the training splits; "
         "re-run the same test grid", "PENDING"],
        ["4", "Recall improvement", "Sep 2026",
         "Self-consistency voting and a rule-based fallback for missed fields; "
         "per-field error analysis", "PENDING"],
        ["5", "Write-up and viva", "Oct 2026 – Nov 2026",
         "Final dissertation, review, and viva", "PENDING"],
    ]
    _table(doc, ["Sl. No", "Phase", "Start – End", "Work to be done", "Status"], rows,
           number=3, caption="Future plan",
           widths=[0.5, 1.5, 1.3, 2.6, 0.9])
    _para(doc,
          "The central activity of the final phase is the fine-tuning in row 3. The aim is to "
          "test whether a small model, tuned on the training data for this narrow task, can "
          "match or beat the large cloud models at a fraction of their cost and latency. The "
          "fine-tuned models are exported back into the same local runner and measured on the "
          "same frozen test set, so the gain over the present baseline can be read off directly.")


def section_abbreviations(doc):
    _heading(doc, "6. Abbreviations")
    rows = [
        ["API", "Application Programming Interface"],
        ["DO", "DigitalOcean"],
        ["F1", "Harmonic mean of precision and recall"],
        ["GGUF", "GPT-Generated Unified Format (model file format for local inference)"],
        ["JSON", "JavaScript Object Notation"],
        ["KIE", "Key Information Extraction"],
        ["LLM", "Large Language Model"],
        ["MoE", "Mixture of Experts"],
        ["OCR", "Optical Character Recognition"],
        ["P / R", "Precision / Recall"],
        ["QAT", "Quantization-Aware Training"],
        ["QLoRA", "Quantized Low-Rank Adaptation (fine-tuning method)"],
        ["RSS", "Resident Set Size (memory in use)"],
        ["SLM", "Small Language Model"],
        ["SROIE", "Scanned Receipts OCR and Information Extraction (dataset)"],
    ]
    _table(doc, ["Abbreviation", "Expansion"], rows, widths=[1.4, 5.0])


# --------------------------------------------------------------------------- main
def main() -> None:
    cfg = load_config()
    rep = cfg["report"]
    plots_dir = resolve_path(cfg, cfg["paths"]["plots_dir"])
    report_dir = resolve_path(cfg, cfg["paths"]["report_dir"])
    fig_dir = report_dir / "figures"

    df = load_results_df(cfg)
    # Table 2 reports the few-shot setting so every column (F1, P, R, exact,
    # latency, cost) comes from the same runs and reconciles. Zero-shot F1 is
    # carried alongside as a labelled reference column for the shot-mode contrast.
    pm = per_model(df[df["shot_mode"] == "few_shot"], cfg)
    zero_f1 = per_model(df, cfg).set_index("model_id")["f1_zero_shot"]
    pm = pm.copy()
    pm["f1_zero_shot"] = pm["model_id"].map(zero_f1)

    doc = Document()
    _set_base_style(doc)

    title_page(doc, rep)
    abstract(doc, rep)
    contents(doc)
    section_modules(doc, fig_dir)
    section_pipeline(doc, fig_dir)
    section_specs(doc, plots_dir, pm)
    section_design(doc)
    section_future(doc)
    section_abbreviations(doc)

    out = report_dir / "Mid_Sem_Report.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
