"""Generate the BITS Mid-Sem report as .docx, mirroring the institute sample
(section order, abstract+signature block, contents, numbered sections, future-plan
and abbreviations tables) — excluding the sample's first page. Project content,
the architecture/pipeline figures, and (when a run exists) the result plots and
per-model table are embedded.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from src.analysis.report_figures import make_figures
from src.config import resolve_path

ABBREVIATIONS = [
    ("API", "Application Programming Interface"),
    ("CSV", "Comma-Separated Values"),
    ("F1", "Harmonic mean of precision and recall"),
    ("GT", "Ground Truth"),
    ("JSON", "JavaScript Object Notation"),
    ("KIE", "Key Information Extraction"),
    ("KV", "Key–Value (pair)"),
    ("LLM", "Large Language Model"),
    ("OCR", "Optical Character Recognition"),
    ("RAM", "Random Access Memory"),
    ("SLM", "Small Language Model"),
    ("SROIE", "Scanned Receipts OCR and Information Extraction (dataset)"),
]

# (Phase, Start–End, Work, Status) — absolute dates; status reflects build state.
FUTURE_PLAN = [
    ("Literature review & proposal", "Jan 2026 – Mar 2026",
     "Survey SLM/LLM KIE work; finalise proposal & deck", "COMPLETED"),
    ("Framework design & build", "Apr 2026 – Jun 2026",
     "Implement data, models, eval, runner, analysis; unit tests", "COMPLETED"),
    ("Pilot run & validation", "Jun 2026",
     "Validate end-to-end pipeline on a small subset", "IN PROGRESS"),
    ("Full evaluation", "Jul 2026",
     "8 models × 2 datasets × conditions on real SROIE/Kleister", "PENDING"),
    ("Analysis, report & viva", "Aug 2026",
     "Frontier analysis, results chapter, final report & demo", "PENDING"),
]


def _center(p):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def _add_toc(document) -> None:
    """Insert a Word TOC field (populates on 'Update Field' in Word)."""
    p = document.add_paragraph()
    run = p.add_run()
    fldchar = OxmlElement("w:fldChar")
    fldchar.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    fldchar_sep = OxmlElement("w:fldChar")
    fldchar_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click → Update Field to build the table of contents."
    fldchar_end = OxmlElement("w:fldChar")
    fldchar_end.set(qn("w:fldCharType"), "end")
    for el in (fldchar, instr, fldchar_sep, placeholder, fldchar_end):
        run._r.append(el)


def _table(document, header: list[str], rows: list[list[str]]):
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for i, h in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
    for r in rows:
        cells = table.add_row().cells
        for i, val in enumerate(r):
            cells[i].text = str(val)
    return table


# ---------------------------------------------------------------------------


def _title_page(document, rep: dict[str, Any]) -> None:
    document.add_paragraph()
    h = _center(document.add_paragraph())
    run = h.add_run(rep["title"].upper())
    run.bold = True
    run.font.size = Pt(18)
    for text in [
        f"{rep['course_code']}: Dissertation", "", "by",
        rep["student_name"], rep["student_id"], "",
        "Dissertation work carried out at", rep["organization"], "",
        f"Submitted in partial fulfilment of {rep['degree_programme']} degree programme", "",
        "Under the Supervision of", rep["supervisor_name"], rep["organization"], "", "",
        rep["institute"], rep["month_year"],
    ]:
        _center(document.add_paragraph(text))
    document.add_page_break()


def _abstract_page(document, synthetic: bool) -> None:
    document.add_heading("ABSTRACT", level=1)
    document.add_paragraph(
        "This dissertation investigates whether small, locally-run language models "
        "(0.36B–7B parameters, served via Ollama) can match large API models (Azure "
        "OpenAI) at Key Information Extraction (KIE) from financial documents, and at "
        "what cost, latency and memory. The same documents — preprocessed into an "
        "AWS-Textract-style \"simple JSON\" of lines, key/value pairs and tables — are "
        "presented to every model, which must return the target fields as JSON."
    )
    document.add_paragraph(
        "A reproducible framework evaluates eight models on two datasets (SROIE "
        "receipts, 4 fields; Kleister-Charity reports, 8 fields) across zero-/few-shot "
        "and input-variant conditions. Extraction quality is scored with precision, "
        "recall, F1 and exact-match (with type-aware normalisation of strings, "
        "currency and dates), while efficiency is measured as latency, peak memory and "
        "USD cost. The output is an accuracy–cost frontier and deployment guidance on "
        "where small models stop being viable as documents grow more complex."
    )
    if synthetic:
        document.add_paragraph(
            "Note: figures in this interim report are produced on a small synthetic "
            "sample used to validate the pipeline; they will be replaced by results on "
            "the real datasets for the final submission."
        )
    document.add_paragraph()
    document.add_paragraph()
    table = document.add_table(rows=1, cols=2)
    left, right = table.rows[0].cells
    left.paragraphs[0].add_run("Signature of the Student").bold = True
    right.paragraphs[0].add_run("Signature of the Supervisor").bold = True
    for cell in (left, right):
        cell.add_paragraph("Name:")
        cell.add_paragraph("Date:")
        cell.add_paragraph("Place:")
    document.add_page_break()


def _contents_page(document) -> None:
    document.add_heading("Contents", level=1)
    _add_toc(document)
    document.add_page_break()


def _section_modules(document, fig1: Path) -> None:
    document.add_heading("1. Framework Modules", level=1)
    document.add_paragraph(
        "The framework is organised into composable modules driven entirely by a single "
        "configuration file (no magic constants in code). Adding a model or dataset "
        "requires only configuration and a loader — never a change to the runner."
    )
    for name, desc in [
        ("Data layer (loaders, ground truth)",
         "Pluggable dataset loaders and a reproducible 20-document test-set selector "
         "(fixed by seed), plus gold-field alignment to each dataset's schema."),
        ("Preprocess",
         "Two converters — one for dataset OCR records, one for raw AWS Textract Blocks "
         "— that emit an identical, deterministic simple JSON in reading order."),
        ("Prompt builder",
         "Builds the model prompt from the target schema and the simple JSON, supporting "
         "zero-/few-shot and lines-only / lines+key-values input variants."),
        ("Model registry & runners",
         "Local models via the Ollama REST API and API models via Azure OpenAI, behind a "
         "common ModelRunner interface; built from configuration."),
        ("Evaluation",
         "Robust JSON parsing with schema validation, type-aware metrics "
         "(precision/recall/F1, exact-match, missing vs hallucinated), and efficiency "
         "measurement (latency, peak memory, cost)."),
        ("Runner & analysis",
         "A resumable grid runner writing incremental results, and an analysis stage "
         "producing summary tables, plots and an auto-written report."),
    ]:
        p = document.add_paragraph(style="List Bullet")
        p.add_run(f"{name}: ").bold = True
        p.add_run(desc)
    document.add_paragraph()
    document.add_picture(str(fig1), width=Inches(6.2))
    _center(document.add_paragraph()).add_run("Figure 1: Framework modules").italic = True


def _section_pipeline(document, fig2: Path) -> None:
    document.add_heading("2. Functional Pipeline", level=1)
    document.add_paragraph(
        "A document is converted (by AWS Textract or dataset OCR) into the simple JSON "
        "representation. The prompt builder combines this with the target schema; the "
        "selected model returns JSON, which is parsed and validated, compared against the "
        "ground truth, and scored. Quality and efficiency measurements accumulate into "
        "per-model and per-condition tables, plots and the final report."
    )
    document.add_picture(str(fig2), width=Inches(6.2))
    _center(document.add_paragraph()).add_run("Figure 2: Functional pipeline").italic = True


def _section_specs(document, cfg: dict[str, Any], results: dict | None) -> None:
    document.add_heading("3. Major Technical Specifications", level=1)
    local = [m for m in cfg["models"] if m["type"] == "local"]
    api = [m for m in cfg["models"] if m["type"] == "api"]
    rows = [
        ["Local models (Ollama)", ", ".join(m["id"] for m in local)],
        ["API models (Azure OpenAI)", ", ".join(m["id"] for m in api)],
        ["Parameter range", "0.36B → 7B (local) + large API models"],
        ["Datasets", "SROIE (4 fields), Kleister-Charity (8 fields)"],
        ["Conditions", "zero-shot / few-shot × lines-only / lines+key-values"],
        ["Test set", f"{cfg['test_set']['total']} documents (reproducible, seeded)"],
        ["Quality metrics", "Precision, Recall, F1 (micro & macro), exact-match"],
        ["Efficiency metrics", "Latency (s/doc), throughput, peak memory (MB), cost (USD)"],
        ["Hardware (local)", "Apple M1, 8 GB unified memory (sequential, keep_alive=0)"],
        ["API provider", cfg["api"]["provider"]],
        ["Reproducibility", f"seed={cfg.get('seed')}, per-run manifest, pinned requirements"],
    ]
    _table(document, ["Parameter", "Specification"], rows)
    _center(document.add_paragraph()).add_run("Table 1: Major technical specifications").italic = True

    if results is not None:
        _embed_results(document, results)


def _embed_results(document, results: dict) -> None:
    document.add_heading("3.1 Preliminary Results", level=2)
    pm = results.get("per_model")
    if pm is not None and len(pm):
        document.add_paragraph(
            "Per-model summary across all conditions (preliminary — see the abstract note "
            "if produced on the synthetic sample):"
        )
        cols = ["model_id", "model_type", "f1_macro", "exact_match_rate",
                "latency_s", "cost_usd_per_doc", "parse_fail_rate"]
        cols = [c for c in cols if c in pm.columns]
        rows = [[("" if v is None else str(v)) for v in (row[c] for c in cols)]
                for _, row in pm.iterrows()]
        _table(document, cols, rows)
        _center(document.add_paragraph()).add_run("Table 2: Per-model results summary").italic = True

    plots = results.get("plots", {})
    fig_no = 3
    captions = {
        "f1_vs_params": "F1 vs model size",
        "cost_frontier": "Accuracy–cost frontier",
        "latency_vs_params": "Latency vs model size",
        "field_heatmap": "Per-field accuracy by model",
    }
    for key, caption in captions.items():
        p = plots.get(key)
        if p and Path(p).exists():
            document.add_picture(str(p), width=Inches(5.8))
            _center(document.add_paragraph()).add_run(f"Figure {fig_no}: {caption}").italic = True
            fig_no += 1


def _section_design(document) -> None:
    document.add_heading("4. Design Considerations", level=1)
    for item in [
        "Reproducibility: fixed seeds, a per-run manifest (config snapshot, model list, "
        "platform), pinned dependency versions, and a seeded 20-document test set.",
        "Pluggability: datasets and models are added via configuration and a loader; the "
        "runner, metrics and analysis are unchanged.",
        "8 GB memory strategy: local models run strictly one at a time, ordered "
        "smallest→largest, unloaded after use (keep_alive=0); the 7B model is opt-in "
        "(--skip-large) so an out-of-memory event never blocks smaller-model results.",
        "Cost control (≈US$25 budget): the six local models are free; only the two Azure "
        "deployments are billed; per-row cost is logged and a pilot validates wiring for "
        "cents before any full run.",
        "Robustness: model output is parsed defensively (markdown stripped, first balanced "
        "JSON object, schema-validated); a parse failure is recorded, never fatal. The run "
        "is resumable — completed cells are skipped.",
        "Reproducible demo: a replay mode renders saved results so the live viva "
        "demonstration cannot fail on network or model-load issues.",
    ]:
        document.add_paragraph(item, style="List Bullet")


def _section_future(document) -> None:
    document.add_heading("5. Future Plan", level=1)
    rows = [[ph, dates, work, status] for (ph, dates, work, status) in FUTURE_PLAN]
    _table(document, ["Phase", "Start – End", "Work to be done", "Status"], rows)


def _section_abbreviations(document) -> None:
    document.add_heading("6. Abbreviations", level=1)
    _table(document, ["Abbreviation", "Expansion"], [list(a) for a in ABBREVIATIONS])


# ---------------------------------------------------------------------------


def build_docx(cfg: dict[str, Any], results: dict | None = None,
               synthetic: bool = True) -> str:
    rep = cfg["report"]
    report_dir = resolve_path(cfg, cfg["paths"]["report_dir"])
    figures = make_figures(report_dir / "figures")

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    _title_page(document, rep)
    _abstract_page(document, synthetic)
    _contents_page(document)
    _section_modules(document, figures["architecture"])
    _section_pipeline(document, figures["pipeline"])
    _section_specs(document, cfg, results)
    _section_design(document)
    _section_future(document)
    _section_abbreviations(document)

    out = report_dir / "MidSem_Report.docx"
    document.save(str(out))
    return str(out)
